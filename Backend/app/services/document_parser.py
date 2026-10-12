import fitz  # PyMuPDF
import os
import re
import logging
import docx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {".txt", ".csv", ".py", ".java", ".js", ".tsx", ".ts", ".html", ".css", ".json"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".md"} | TEXT_EXTENSIONS | IMAGE_EXTENSIONS


class DocumentParseError(Exception):
    """Text extraction failed. str(e) is a short, user-safe reason."""


class UnsupportedFileTypeError(DocumentParseError):
    """The file extension has no extractor."""


def extract_text_from_pdf(file_path: str) -> str:
    """Extracts raw text from a PDF file using PyMuPDF."""
    text = ""
    try:
        doc = fitz.open(file_path)
        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            text += page.get_text("text") + "\n\n"
        doc.close()
    except Exception as e:
        logger.exception("PDF parsing failed")
        raise DocumentParseError("Could not read this PDF file") from e
    return text.strip()

def extract_text_from_docx(file_path: str) -> str:
    """Extracts text from a DOCX file, mapping Word heading styles to Markdown headings."""
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        logger.exception("DOCX parsing failed")
        raise DocumentParseError("Could not read this Word document") from e

    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        level = re.match(r"Heading (\d)", style)
        if style == "Title":
            lines.append(f"# {text}")
        elif level:
            lines.append(f"{'#' * min(int(level.group(1)) + 1, 6)} {text}")
        else:
            lines.append(text)
    return "\n\n".join(lines)


def _iter_shapes(shapes):
    """Yields shapes, recursing into group shapes."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def extract_text_from_pptx(file_path: str) -> str:
    """Extracts slide titles, body text, tables and speaker notes from a PPTX file."""
    try:
        prs = Presentation(file_path)
    except Exception as e:
        logger.exception("PPTX parsing failed")
        raise DocumentParseError("Could not read this PowerPoint file") from e

    sections = []
    for index, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        parts = [f"## Slide {index}: {title}" if title else f"## Slide {index}"]

        for shape in _iter_shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [
                    "| " + " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells) + " |"
                    for row in shape.table.rows
                ]
                if rows:
                    col_count = len(shape.table.columns)
                    rows.insert(1, "|" + " --- |" * col_count)
                    parts.append("\n".join(rows))

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                parts.append(f"> **Speaker notes:** {notes}")

        sections.append("\n\n".join(parts))

    return "\n\n".join(sections).strip()

def process_markdown_frontmatter(text: str) -> str:
    """
    Detects YAML frontmatter in a Markdown string and converts it to a
    proper Markdown table so it renders beautifully in the viewer.
    Handles both --- delimited and bare YAML-at-top formats.
    """
    # Normalize line endings
    text = text.replace('\r\n', '\n').replace('\r', '\n').strip()

    # --- Case 1: Standard YAML frontmatter with --- delimiters ---
    fm_match = re.match(r'^-{3,}\s*\n([\s\S]*?)\n-{3,}\s*\n?', text)
    yaml_block = None
    body = text

    if fm_match:
        yaml_block = fm_match.group(1)
        body = text[fm_match.end():]
    else:
        # --- Case 2: Detect bare YAML-like lines at the top ---
        lines = text.split('\n')
        yaml_end = -1
        yaml_line_re = re.compile(r'^[a-zA-Z_][a-zA-Z0-9_]*:\s*.+')
        for i, line in enumerate(lines[:20]):
            stripped = line.strip()
            if stripped == '':
                if yaml_end >= 0:
                    break
                continue
            if yaml_line_re.match(line) or line.startswith(('  ', '\t')):
                yaml_end = i
            else:
                if i == 0:
                    return text  # First line is not YAML, do nothing
                break
        if yaml_end >= 1:
            yaml_block = '\n'.join(lines[:yaml_end + 1])
            body = '\n'.join(lines[yaml_end + 1:]).lstrip('\n')

    if not yaml_block:
        return text

    # Parse the YAML block into key-value pairs
    meta = {}
    current_key = None
    current_value = []

    for line in yaml_block.split('\n'):
        if not line.strip():
            continue
        top_level = re.match(r'^([a-zA-Z_][a-zA-Z0-9_]*):\s*(.*)', line)
        if top_level:
            if current_key:
                val = ' '.join(current_value).strip().strip('"\'')
                meta[current_key] = val
            current_key = top_level.group(1)
            current_value = [top_level.group(2)]
        else:
            current_value.append(line.strip())

    if current_key:
        val = ' '.join(current_value).strip().strip('"\'')
        meta[current_key] = val

    if not meta:
        return text

    # Build the Markdown table
    rows = '\n'.join(
        f"| **{k.replace('_', ' ').title()}** | {v} |"
        for k, v in meta.items() if v.strip()
    )
    table = f"| Property | Value |\n| :--- | :--- |\n{rows}\n\n"

    return table + body


def extract_text_from_file(file_path: str) -> str:
    """Universal text extractor based on file extension."""
    """Universal text extractor based on file extension.

    Raises DocumentParseError (or UnsupportedFileTypeError) on failure; never
    returns an error message as content. Images yield an empty string.
    """
    if not os.path.exists(file_path):
        raise DocumentParseError("Uploaded file could not be found")

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext == ".pptx":
        return extract_text_from_pptx(file_path)
    elif ext == ".md":
        # For Markdown files, process frontmatter server-side for clean rendering
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                raw = f.read()
            return process_markdown_frontmatter(raw)
        except UnicodeDecodeError as e:
            raise DocumentParseError("File is not valid UTF-8 text") from e
    elif ext in TEXT_EXTENSIONS:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        except UnicodeDecodeError as e:
            raise DocumentParseError("File is not valid UTF-8 text") from e
    elif ext in IMAGE_EXTENSIONS:
        # No text extraction for images; never expose the server path as content.
        return ""
    else:
        raise UnsupportedFileTypeError(f"Unsupported file type: {ext or 'no extension'}")

